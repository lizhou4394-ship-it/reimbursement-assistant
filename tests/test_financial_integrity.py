"""用户视角的日期证据、导出合计及输入异常回归。"""
import unittest
from io import BytesIO
from openpyxl import load_workbook
from config import get_builtin_template_bytes
from services.excel_generator import ExcelGenerator
from services.record_validation import normalize_record, normalize_amount
from services.invoice_parser import InvoiceParser
from tests.test_parser_integrity import train_text


class FinancialIntegrityTest(unittest.TestCase):
    def generator(self):
        return ExcelGenerator(get_builtin_template_bytes())

    def test_disjoint_hotel_stays_in_same_city_are_not_lost(self):
        records = [{"type": "酒店", "date": "2026-09-12", "hotel_name": "温州市酒店", "amount": 300,
                    "check_in_date": start, "check_out_date": end}
                   for start, end in [("2026-09-01", "2026-09-02"), ("2026-09-10", "2026-09-11")]]
        self.assertEqual(self.generator().calculate_travel_days(records), {"温州": 4})

    def test_overlapping_hotel_stays_count_dates_once(self):
        records = [{"type": "酒店", "date": "2026-09-04", "hotel_name": "温州市酒店", "amount": 300,
                    "check_in_date": start, "check_out_date": end}
                   for start, end in [("2026-09-01", "2026-09-03"), ("2026-09-02", "2026-09-04")]]
        self.assertEqual(self.generator().calculate_travel_days(records), {"温州": 4})

    def test_hotel_and_transport_disjoint_days_are_combined(self):
        records = [{"type": "火车票", "date": "2026-09-01", "start_location": "杭州东", "end_location": "温州南"},
                   {"type": "火车票", "date": "2026-09-02", "start_location": "温州南", "end_location": "杭州东"},
                   {"type": "酒店", "date": "2026-09-12", "hotel_name": "温州市酒店", "amount": 300,
                    "check_in_date": "2026-09-10", "check_out_date": "2026-09-11"}]
        self.assertEqual(self.generator().calculate_travel_days(records), {"温州": 4})

    def test_substitute_remark_survives_g_column_subtotal(self):
        records = [{"type": "打车", "date": "2026-09-01", "start_location": "杭州东", "end_location": "客户",
                    "amount": 30, "need_substitute": True, "itinerary_file": "a.pdf"}]
        ws = load_workbook(BytesIO(self.generator().generate(records, {}, ""))).active
        self.assertEqual(ws["G3"].value, 30)
        self.assertIn("餐费代替", ws["D3"].value or "")

    def test_mixed_export_totals_include_only_entertainment_meals(self):
        records = [{"type": kind, "date": "2026-09-01", "amount": amount, **extra}
                   for kind, amount, extra in [
                       ("火车票", 100, {"start_location": "杭州东", "end_location": "温州南"}),
                       ("酒店", 200, {"hotel_name": "温州市酒店"}),
                       ("餐饮", 300, {}), ("餐饮", 50, {"is_entertainment": True}),
                       ("快递", 20, {}), ("其他", 10, {})]]
        ws = load_workbook(BytesIO(self.generator().generate(records, {"温州": 2}, ""))).active
        self.assertEqual(next(row[5].value for row in ws.iter_rows() if row[0].value == "总合计"), 1180)

    def test_giant_numbers_cannot_overflow_to_infinity(self):
        with self.assertRaises(ValueError):
            normalize_amount("9" * 400)

    def test_boolean_flags_are_not_truthy_strings(self):
        record = normalize_record({"type": "餐饮", "date": "2026-09-01", "amount": 100, "is_entertainment": "false"})
        self.assertIs(record["is_entertainment"], False)

    def test_invalid_hotel_dates_nights_and_rates_are_rejected(self):
        base = {"type": "酒店", "date": "2026-09-01", "amount": 300}
        for change in [{"check_in_date": "2026-09-03", "check_out_date": "2026-09-01"},
                       {"check_in_date": "2026-09-01"}, {"nights": "1.5"}, {"daily_rate": "NaN"}]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                normalize_record({**base, **change})

    def test_travel_date_wins_even_when_invoice_date_has_time(self):
        result = InvoiceParser("")._try_parse_regex(train_text(extra="开票日期：2026-10-07 16:30"))
        self.assertEqual(result["date"], "2026-09-05")

    def test_labeled_unknown_stations_override_document_order(self):
        text = train_text().replace("杭州西站\n衢州西站", "到达站：峨眉山站\n出发站：成都东站")
        result = InvoiceParser("")._try_parse_regex(text)
        self.assertEqual((result["start_location"], result["end_location"]), ("成都东", "峨眉山"))

    def test_transfer_day_is_paid_once_and_city_details_are_preserved(self):
        records = [{"type": "火车票", "date": day, "start_location": start, "end_location": end, "amount": 100}
                   for day, start, end in [("2026-09-01", "杭州东", "温州南"),
                                           ("2026-09-02", "温州南", "永康南"),
                                           ("2026-09-02", "永康南", "杭州东")]]
        generator = self.generator()
        days = generator.calculate_travel_days(records)
        self.assertEqual(days, {"温州": 2, "永康": 1})
        self.assertEqual(generator.total_travel_days(days), 2)
        ws = load_workbook(BytesIO(generator.generate(records, days, ""))).active
        self.assertEqual(next(row[5].value for row in ws.iter_rows() if row[4].value == "出差补贴"), 100)
        self.assertEqual(next(row[5].value for row in ws.iter_rows() if row[0].value == "总合计"), 1100)

    def test_home_city_taxi_after_return_does_not_extend_trip(self):
        records = [{"type": "火车票", "date": "2026-09-01", "start_location": "杭州东", "end_location": "温州南"},
                   {"type": "火车票", "date": "2026-09-02", "start_location": "温州南", "end_location": "杭州东"},
                   {"type": "打车", "date": "2026-09-03", "start_location": "杭州市西湖区", "end_location": "杭州市公司"}]
        self.assertEqual(self.generator().calculate_travel_days(records), {"温州": 2})
