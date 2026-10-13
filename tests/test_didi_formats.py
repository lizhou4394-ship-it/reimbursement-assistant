import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook
from services.data_correlator import DataCorrelator
from services.excel_generator import ExcelGenerator
from services.invoice_parser import InvoiceParser
from test_excel_generator import blank_template_bytes


def document(rows, total="30.00", period="2026-09-01 至 2026-09-30"):
    return f"滴滴出行-行程单\n申请日期：2027-01-05\n行程起止日期：{period}\n共{len(rows)}笔行程，合计{total}元\n备注\n" + "\n".join(rows)


def row(index=1, day="09-05", clock="14:18", car="特惠快车", amount="30.00", distance="12.0", origin="区域|市一医院", destination="街道|火车东站"):
    return f"{index}\n{car}\n{day} {clock} 周\n一\n杭州\n市\n{origin}\n{destination}\n{distance}\n{amount}\n"


class DidiFormatsTest(unittest.TestCase):
    def setUp(self):
        self.parser = InvoiceParser("")

    def test_whitespace_fullwidth_and_time_variants(self):
        for clock in ["14:\n18", "14\n:\n18", "14 : 18", "１４：１８"]:
            with self.subTest(clock=clock):
                text = document([row(clock=clock)]).replace("共1笔", "共 1 笔").replace("滴滴出行", "滴滴 出行")
                result = self.parser._try_parse_regex(text)
                self.assertEqual(result[0]["time"], "14:18")
                self.assertEqual(result[0]["date"], "2026-09-05")

    def test_distance_and_address_decimals_do_not_become_amounts(self):
        result = self.parser._try_parse_regex(document([row(distance="12.34", origin="区域|30.50便利店")]))
        self.assertEqual(result[0]["amount"], 30)

    def test_large_and_one_decimal_amounts(self):
        for amount in ["500.00", "1,200.50", "30.5"]:
            with self.subTest(amount=amount):
                result = self.parser._try_parse_regex(document([row(amount=amount)], total=amount))
                self.assertEqual(result[0]["amount"], float(amount.replace(",", "")))

    def test_car_types_are_bound_to_their_rows(self):
        text = document([
            row(car="新\n车型", origin="区域|顺风车服务站"),
            row(2, car="顺\n风\n车"), row(3, car="特惠\n快车"),
        ], total="90.00")
        result = self.parser._try_parse_regex(text)
        self.assertEqual([r["car_type"] for r in result], ["新车型", "顺风车", "特惠快车"])
        self.assertEqual([r["need_substitute"] for r in result], [False, True, False])

    def test_cross_year_and_application_year(self):
        text = document([row(day="12-31"), row(2, day="01-01")], total="60.00", period="2025-12-31 至 2026-01-01")
        result = self.parser._try_parse_regex(text)
        self.assertEqual([r["date"] for r in result], ["2025-12-31", "2026-01-01"])

    def test_bad_calendar_time_count_total_sequence_and_missing_year_rejected(self):
        base = document([row()])
        variants = [base.replace("09-05", "02-30"), base.replace("14:18", "25:18"),
                    base.replace("共1笔", "共2笔"), base.replace("合计30.00", "合计31.00"),
                    base.replace("备注\n1\n", "备注\n2\n"), base.replace("行程起止日期", "其他日期")]
        for text in variants:
            with self.subTest(text=text):
                self.assertIsNone(self.parser._try_parse_regex(text))

    def test_city_prefix_does_not_remove_first_character_of_origin(self):
        result = self.parser._try_parse_regex(document([row(origin="市一医院-停车场")]))
        self.assertEqual(result[0]["start_location"], "市一医院-停车场")

    def test_wrapped_parenthesis_belongs_to_destination(self):
        result = self.parser._try_parse_regex(document([row(origin="区域|博览中心E5门", destination="市心北路汉庭(博览中心飞虹路地\n铁站店)")]))
        self.assertEqual(result[0]["start_location"], "博览中心E5门")
        self.assertEqual(result[0]["end_location"], "市心北路汉庭")

    def test_26_rows_with_repeated_headers_and_wrapped_times(self):
        amounts = [50.5,38.1,46.4,43.4,42.8,12.4,12.4,12.4,12.4,21.3,29.4,52.1,15.3,12,15,28,15.3,42.7,54.7,35.8,63.3,45.5,45.3,23.6,15.3,32.4]
        rows = [row(i+1, day=f"09-{i+5:02d}", clock="16:\n17" if 10 <= i < 25 else "14:18", amount=f"{amount:.2f}") for i,amount in enumerate(amounts)]
        for index in [9,24]:
            rows[index] += "页码：2/3\n序号\n车型\n上车时间\n城市\n起点\n终点\n里程[公里] 金额[元] 备注\n"
        result = self.parser._try_parse_regex(document(rows, total="817.80"))
        self.assertEqual(len(result), 26)
        self.assertAlmostEqual(sum(r["amount"] for r in result), 817.8)
        for record in result:
            record["source_file"] = "sample.pdf"
            record["source_file_rank"] = 0
        correlated = DataCorrelator("").correlate_all(result, "")
        ws = load_workbook(BytesIO(ExcelGenerator(blank_template_bytes()).generate(correlated, {}, ""))).active
        self.assertEqual([ws.cell(r, 6).value for r in range(3, 29)], amounts)
        self.assertEqual(ws["G3"].value, 817.8)
        self.assertIn("G3:G28", [str(rng) for rng in ws.merged_cells.ranges])
        subtotal = next(r for r in ws.iter_rows() if r[0].value == "城际交通 小计")
        self.assertEqual(subtotal[5].value, 817.8)

    def test_ai_dates_time_and_fields_are_checked(self):
        text = document([row()])
        original = self.parser._try_parse_regex(text)[0]
        for key, value in [("date", "2027-09-05"), ("date", "2026-02-30"),
                           ("time", "25:18"), ("time", ""), ("car_type", ""),
                           ("amount", float("nan")), ("amount", 31)]:
            with self.subTest(key=key, value=value):
                self.assertFalse(self.parser._valid_didi_result(text, {**original, key: value}))
        self.assertTrue(self.parser._valid_didi_result(text, original))

    def test_partial_text_extraction_uses_all_images(self):
        text = document([row(), row(2)], total="60.00")
        records = self.parser._try_parse_regex(text)
        with patch.object(self.parser, "_parse_with_vision_model", side_effect=[[records[0]], [records[1]]]), patch.object(self.parser, "_parse_with_text_model") as model:
            result = self.parser.parse_all_invoices([{"filename":"partial.pdf", "page_texts":[text], "images":["p1", "p2"]}])
        self.assertEqual(len(result), 2)
        self.assertEqual(self.parser.failed_files, [])
        model.assert_not_called()

    def test_text_fallback_mismatch_uses_all_images(self):
        text = document([row(), row(2)], total="60.00")
        records = self.parser._try_parse_regex(text)
        with patch.object(self.parser, "_try_parse_regex", return_value=None), patch.object(self.parser, "_parse_with_text_model", return_value=records[:1]), patch.object(self.parser, "_parse_with_vision_model", side_effect=[[records[0]], [records[1]]]) as vision:
            result = self.parser.parse_all_invoices([{"filename":"sample.pdf", "page_texts":[text, "continuation"], "images":["p1", "p2"]}])
        self.assertEqual(len(result), 2)
        self.assertEqual([call.args[0] for call in vision.call_args_list], ["p1", "p2"])
        self.assertEqual([r["trip_index"] for r in result], [1, 2])

    def test_missing_scanned_page_is_reported_instead_of_partial_success(self):
        text = document([row(), row(2)], total="60.00")
        records = self.parser._try_parse_regex(text)
        for responses in [[records[:1], None], [records[:1], []]]:
            with self.subTest(responses=responses), patch.object(self.parser, "_parse_with_vision_model", side_effect=responses), patch.object(self.parser, "_parse_with_text_model") as model:
                result = self.parser.parse_all_invoices([{"filename":"mixed.pdf", "page_texts":[text, ""], "images":["p1", "p2"]}])
                self.assertEqual(result, [])
                self.assertEqual(self.parser.failed_files, ["mixed.pdf"])
                model.assert_not_called()

    def test_intentionally_skipped_invoice_is_not_reported_as_failed(self):
        result = self.parser.parse_all_invoices([{"filename":"invoice.pdf", "page_texts":["电子发票 旅客运输服务"], "images":[]}])
        self.assertEqual(result, [])
        self.assertEqual(self.parser.failed_files, [])
