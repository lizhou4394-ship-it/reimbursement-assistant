import os
import sys
import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook


sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from services.excel_generator import ExcelGenerator


def blank_template_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "报销单"
    ws["A1"] = "报销单"
    ws["A2"] = "日期"
    ws["B2"] = "起点"
    ws["C2"] = "终点"
    ws["D2"] = "工作内容和目的"
    ws["E2"] = "费用\n名称"
    ws["F2"] = "金额(元)"
    ws["G2"] = "备注(超标原因/替票、替票原因）"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class ExcelGeneratorRegressionTest(unittest.TestCase):
    def test_travel_days_use_transport_segments(self):
        generator = ExcelGenerator(blank_template_bytes())
        invoices = [
            {"type": "火车票", "date": "2026-06-08", "time": "06:35", "start_location": "杭州西", "end_location": "温州北"},
            {"type": "火车票", "date": "2026-06-10", "time": "13:55", "start_location": "温州南", "end_location": "永康南"},
            {"type": "火车票", "date": "2026-06-10", "time": "16:29", "start_location": "永康南", "end_location": "杭州东"},
            {"type": "火车票", "date": "2026-06-29", "time": "20:00", "start_location": "浦江", "end_location": "杭州西"},
            {"type": "火车票", "date": "2026-07-05", "time": "18:00", "start_location": "杭州西", "end_location": "温州南"},
            {"type": "打车", "date": "2026-07-06", "start_location": "桔子酒店", "end_location": "温州医科大学附属第一医院"},
        ]

        days = generator.calculate_travel_days(invoices)

        self.assertEqual(days["温州"], 5)
        self.assertEqual(days["永康"], 1)
        self.assertEqual(days["浦江"], 1)

    def test_travel_days_include_both_departure_and_return_dates(self):
        generator = ExcelGenerator(blank_template_bytes())
        invoices = [
            {"type": "火车票", "date": "2026-09-01", "time": "23:00", "start_location": "杭州东", "end_location": "温州南"},
            {"type": "火车票", "date": "2026-09-02", "time": "01:00", "start_location": "温州南", "end_location": "杭州东"},
        ]

        self.assertEqual(generator.calculate_travel_days(invoices), {"温州": 2})

    def test_taxi_date_fills_missing_transport_day(self):
        generator = ExcelGenerator(blank_template_bytes())
        invoices = [
            {"type": "火车票", "date": "2026-09-01", "time": "23:00", "start_location": "杭州东", "end_location": "温州南"},
            {"type": "打车", "date": "2026-09-02", "time": "01:00", "start_location": "温州南", "end_location": "温州酒店"},
        ]

        self.assertEqual(generator.calculate_travel_days(invoices), {"温州": 2})

    def test_explicit_hotel_stay_uses_inclusive_natural_days(self):
        generator = ExcelGenerator(blank_template_bytes())
        invoices = [
            {
                "type": "酒店",
                "date": "2026-09-03",
                "check_in_date": "2026-09-01",
                "check_out_date": "2026-09-02",
                "hotel_name": "温州某酒店",
                "amount": 300,
            },
            {
                "type": "火车票",
                "date": "2026-09-01",
                "start_location": "杭州东",
                "end_location": "温州南",
            },
            {
                "type": "打车",
                "date": "2026-09-01",
                "start_location": "温州南",
                "end_location": "温州某酒店",
            },
        ]

        self.assertEqual(generator.calculate_travel_days(invoices), {"温州": 2})

    def test_taxi_rows_are_grouped_by_itinerary_file_and_g_has_group_total(self):
        generator = ExcelGenerator(blank_template_bytes())
        invoices = [
            {"type": "打车", "date": "2026-07-10", "time": "10:00", "trip_index": 2, "source_file_rank": 0, "itinerary_rank": 0, "itinerary_file": "itinerary-a.pdf", "start_location": "温州南", "end_location": "医院", "amount": 20},
            {"type": "打车", "date": "2026-07-01", "time": "09:00", "trip_index": 1, "source_file_rank": 0, "itinerary_rank": 0, "itinerary_file": "itinerary-a.pdf", "start_location": "温州南", "end_location": "酒店", "amount": 30},
            {"type": "打车", "date": "2026-07-02", "time": "08:00", "trip_index": 1, "source_file_rank": 1, "itinerary_rank": 1, "itinerary_file": "itinerary-b.pdf", "start_location": "嘉兴南", "end_location": "客户", "amount": 40},
        ]

        output = generator.generate(invoices, {}, "")
        ws = load_workbook(BytesIO(output)).active

        self.assertEqual([ws.cell(r, 1).value for r in range(3, 6)], ["2026.7.1", "2026.7.10", "2026.7.2"])
        self.assertEqual(ws["G3"].value, 50)
        self.assertEqual(ws["G5"].value, 40)
        self.assertTrue(any(str(rng) == "G3:G4" for rng in ws.merged_cells.ranges))


if __name__ == "__main__":
    unittest.main()
