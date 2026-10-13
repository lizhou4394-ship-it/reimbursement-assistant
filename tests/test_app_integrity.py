"""Regression: navigation bypassed validation and retained stale totals, QA 2026-10-13."""
import os
import unittest
from io import BytesIO
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from config import get_builtin_template_bytes


class AppIntegrityTest(unittest.TestCase):
    def app(self):
        app = AppTest.from_file(os.path.join(os.path.dirname(os.path.dirname(__file__)), "app.py"), default_timeout=20)
        app.session_state.invoices = [
            {"type": "火车票", "date": "2026-09-01", "start_location": "杭州东", "end_location": "温州南", "amount": 93},
            {"type": "火车票", "date": "2026-09-02", "start_location": "温州南", "end_location": "杭州东", "amount": 93},
        ]
        app.session_state.current_step = 1
        app.session_state.template_bytes = get_builtin_template_bytes()
        app.session_state.travel_days = {"温州": 2}
        return app.run()

    def export(self, app):
        next(b for b in app.button if b.label == "3️⃣ 生成报销单").click().run()

    def test_invalid_edit_cannot_bypass_validation_via_navigation(self):
        for bad in ["NaN", "Infinity", "1.234", "-10", "abc"]:
            with self.subTest(bad=bad):
                app = self.app()
                app.text_input(key="amount_0").set_value(bad).run()
                self.export(app)
                self.assertTrue(app.error)
                self.assertFalse(app.session_state.generated_excel)
                self.assertFalse(app.exception)

    def test_date_edit_refreshes_days_without_pressing_save(self):
        app = self.app()
        app.text_input(key="date_1").set_value("2026-09-03").run()
        self.export(app)
        self.assertEqual(app.session_state.travel_days, {"温州": 3})
        self.assertTrue(app.session_state.generated_excel)
        self.assertFalse(app.exception)

    def test_changed_data_cannot_offer_old_download_after_generation_failure(self):
        app = self.app()
        self.export(app)
        self.assertTrue(app.session_state.generated_excel)
        next(b for b in app.button if b.label == "2️⃣ 预览与编辑").click().run()
        app.text_input(key="date_0").set_value("").run()
        self.export(app)
        self.assertTrue(app.error)
        self.assertFalse(app.session_state.generated_excel)

    def test_same_event_navigation_preserves_unsubmitted_edit(self):
        app = self.app()
        app.text_input(key="amount_0").set_value("120.50")
        self.export(app)
        ws = load_workbook(BytesIO(app.session_state.generated_excel)).active
        self.assertEqual(ws["F3"].value, 120.5)

    def test_partial_file_failures_require_acknowledgement(self):
        app = self.app()
        app.session_state.failed_files = ["unreadable.pdf"]
        self.export(app)
        self.assertFalse(app.session_state.generated_excel)
        self.assertTrue(app.warning)
        app.checkbox(key="confirm_partial_export").check().run()
        self.assertTrue(app.session_state.generated_excel)

    def test_summary_total_matches_download_total(self):
        app = self.app()
        self.export(app)
        ws = load_workbook(BytesIO(app.session_state.generated_excel)).active
        total = next(row[5].value for row in ws.iter_rows() if row[0].value == "总合计")
        metric = next(m for m in app.metric if m.label == "导出总金额")
        self.assertEqual(metric.value, f"¥{total:.2f}")

    def test_prompt_reset_restores_widget_and_configuration(self):
        from config import load_prompt
        app = self.app()
        app.text_area(key="ta_hotel_infer").set_value("custom prompt").run()
        app.button(key="reset_hotel_infer").click().run()
        self.assertEqual(app.text_area(key="ta_hotel_infer").value, load_prompt("hotel_infer"))
        self.assertEqual(app.session_state.hotel_infer_prompt, load_prompt("hotel_infer"))


if __name__ == "__main__":
    unittest.main()
